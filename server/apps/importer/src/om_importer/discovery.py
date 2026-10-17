"""Read-only structural discovery and a conservative default configuration."""
import re
from docx import Document
from openpyxl import load_workbook
from .common import scalar
from .docx import LABEL, MONTHS
from .xlsx import HEADERS


def inspect_source(path):
    if path.suffix.lower()=='.xlsx':
        w=load_workbook(path,data_only=False)
        try:
            sheets=[]
            for s in w:
                nonempty=next((row for row in s if any(c.value is not None for c in row)),None)
                sheets.append({'name':s.title,'rows':s.max_row,'columns':s.max_column,
                    'suggested_header_row':nonempty[0].row if nonempty else None,
                    'headers':[scalar(c.value) for c in nonempty] if nonempty else []})
            return {'file_type':'xlsx','worksheets':sheets}
        finally:
            w.close()
    d=Document(path)
    return {'file_type':'docx','paragraph_count':len(d.paragraphs),
            'paragraph_examples':[{'index':i,'text':p.text} for i,p in enumerate(d.paragraphs) if p.text.strip()][:8],
            'tables':[{'index':i,'rows':len(t.rows),'headers':[c.text for c in t.rows[0].cells] if t.rows else []} for i,t in enumerate(d.tables)]}


def default_configuration(path):
    info=inspect_source(path)
    if info['file_type']=='xlsx':
        populated=[s for s in info['worksheets'] if s['suggested_header_row'] is not None]
        if not populated:
            raise ValueError('Workbook has no populated worksheets')
        # The sample preset is optional and selected only when its evidence matches.
        if any(s['headers']==HEADERS for s in populated):
            return None
        if len(populated)>1:
            raise ValueError('Multiple populated worksheets; specify sheet and mappings using --config (see inspect)')
        s=populated[0]
        used=set();fields=[]
        from openpyxl.utils import get_column_letter
        for i,header in enumerate(s['headers'],1):
            label=str(header) if header is not None else f'Column {get_column_letter(i)}'
            name=re.sub('[^a-z0-9]+','_',label.lower()).strip('_')[:60]
            if not name or not name[0].isalpha():
                name='field_'+name
            name=name[:60];base=name;suffix=1
            while name in used:
                suffix+=1;name=f'{base}_{suffix}'
            used.add(name)
            fields.append({'name':name,'label':label,'type':'string','source':{'column':get_column_letter(i)}})
        return {'profile':'xlsx','selection':{'sheet':s['name'],'header_row':s['suggested_header_row']},'fields':fields}
    d=Document(path)
    if any(p.text.strip() in MONTHS for p in d.paragraphs) and any(LABEL.match(p.text) for p in d.paragraphs):
        return None
    return {'profile':'docx_paragraphs','fields':[{'name':'text','label':'Text','type':'string','source':{'path':'text'}}]}
