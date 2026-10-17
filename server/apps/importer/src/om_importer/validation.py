import calendar
from functools import lru_cache
from datetime import date
import json
import math
import re
from importlib.resources import files
from jsonschema import Draft202012Validator, FormatChecker


@lru_cache(maxsize=3)
def schema_for(version):
    if version not in ('1.0.0','2.0.0','2.1.0'):
        raise ValueError(f'Unsupported package schema version: {version}')
    return json.loads(files('om_importer').joinpath(f'schemas/import-package-{version}.json').read_text())


def validate_partial_date(d):
    schema=schema_for('1.0.0')
    Draft202012Validator(schema['$defs']['partial_date']).validate(d)
    if d['precision']=='year' and (d['month'] is not None or d['day'] is not None):
        raise ValueError('Year precision cannot carry month/day')
    if d['precision']=='month' and (not d['month'] or d['day'] is not None):
        raise ValueError('Invalid month precision')
    if d['precision']=='day' and (not d['month'] or not d['day']):
        raise ValueError('Day precision requires month/day')
    if d['day'] and d['day']>calendar.monthrange(d['year'],d['month'])[1]:
        raise ValueError('Invalid partial date')


def validate_recurrence(rule):
    schema=schema_for('1.0.0')
    Draft202012Validator({'$defs':schema['$defs'],'$ref':'#/$defs/recurrence'}).validate(rule)
    if rule['type']=='fixed' and rule['day']>calendar.monthrange(2000,rule['month'])[1]:
        raise ValueError('Invalid fixed annual date')
    if rule['type']=='weekday_in_range' and rule['parameters']['start_day']>rule['parameters']['end_day']:
        raise ValueError('Recurrence range is reversed')


def validate_preset(values,profile):
    schema=schema_for('1.0.0')
    definition='parish' if profile=='parish_announcements_v1' else 'cemetery'
    Draft202012Validator({'$defs':schema['$defs'],'$ref':'#/$defs/'+definition}).validate(values)
    if definition=='parish':
        rule=values['recurrence']
        if rule['type']=='fixed' and rule['day']>calendar.monthrange(2000,rule['month'])[1]:
            raise ValueError('Invalid fixed annual date')
        if rule['type']=='weekday_in_range' and rule['parameters']['start_day']>rule['parameters']['end_day']:
            raise ValueError('Recurrence range is reversed')
        if (values['event_year'] is None)!=(values['year_role'] is None):
            raise ValueError('Year and role must be paired')
    else:
        for key in ('birth','death'):
            if values[key] is not None:
                validate_partial_date(values[key])


def validate_typed_value(value,field):
    typ=field['type']
    if value is None:
        return
    if typ=='string' and not isinstance(value,str):
        raise ValueError('Expected string field')
    if typ=='integer' and (type(value) is not int or not -(2**63)<=value<2**63):
        raise ValueError('Expected BIGINT-compatible integer field')
    if typ=='number' and (type(value) not in (int,float) or not math.isfinite(value)):
        raise ValueError('Expected finite number field')
    if typ=='boolean' and type(value) is not bool:
        raise ValueError('Expected boolean field')
    if typ=='date':
        if not isinstance(value,str) or date.fromisoformat(value).isoformat()!=value:
            raise ValueError('Expected ISO full date field')
    if typ=='partial_date':
        validate_partial_date(value)
    if typ=='json':
        json.dumps(value,allow_nan=False)


def validate(p):
    """Check declared schema, mapped types, provenance and review gates without source I/O."""
    version=p.get('schema_version')
    schema=schema_for(version)
    Draft202012Validator(schema,format_checker=FormatChecker()).validate(p)
    json.dumps(p,allow_nan=False)
    config=None
    if version in ('2.0.0','2.1.0'):
        from .configuration import validate_config,field_columns
        config=validate_config(p['configuration'])
        if config['profile']!=p['extraction']['profile']:
            raise ValueError('Configuration/extraction profile mismatch')
        expected_format='xlsx' if config['profile'] in ('xlsx','cemetery_v1') else 'docx'
        if p['source']['file_type']!=expected_format:
            raise ValueError('Configuration/source format mismatch')
        if p['resource']['columns']!=field_columns(config):
            raise ValueError('Column definitions do not match import configuration')
    blocks={b['id']:b for b in p['source_blocks']}
    if len(blocks)!=len(p['source_blocks']):
        raise ValueError('Duplicate source block identifiers')
    bindings = None
    if config and config['profile'] in ('xlsx','docx_table'):
        from .generic import column_bindings
        sel = config['selection']
        header = next((b for b in p['source_blocks'] if b['location'].get('row') == sel.get('header_row')
                       and (b['location'].get('sheet') == sel.get('sheet') if config['profile']=='xlsx'
                            else b['location'].get('table') == sel.get('table_index'))),None)
        if not header:
            raise ValueError('Configured source header block is missing')
        headers = [v['value'] for v in header['raw'].values()] if config['profile']=='xlsx' else list(header['raw'].values())
        bindings = column_bindings(config,headers)
    paragraph_pattern = None
    if config and config['profile']=='docx_paragraphs' and config['selection'].get('pattern'):
        paragraph_pattern = re.compile(config['selection']['pattern'])
    ids=set();issue_ids=set()
    # Validate issue identifiers before using them to explain missing required values.
    record_ids={r['id'] for r in p['records']}
    for i in p['issues']:
        if i['id'] in issue_ids or i['block_id'] not in blocks or (i['record_id'] is not None and i['record_id'] not in record_ids):
            raise ValueError('Invalid issue reference or duplicate issue identifier')
        issue_ids.add(i['id'])
    column_names=[c['name'] for c in p['resource']['columns']]
    if len(column_names)!=len(set(column_names)):
        raise ValueError('Duplicate column names')
    for r in p['records']:
        if r['id'] in ids:
            raise ValueError('Duplicate record identifier')
        ids.add(r['id'])
        b=blocks.get(r['block_id'])
        if not b:
            raise ValueError('Missing source block')
        loc=r['source_location']
        if any(loc.get(k)!=v for k,v in b['location'].items()):
            raise ValueError('Source location disagrees with block')
        if isinstance(b['raw'],str):
            a,z=loc.get('char_start'),loc.get('char_end')
            if a is None or z is None or not 0<=a<=z<=len(b['raw']) or r['original']!=b['raw'][a:z]:
                raise ValueError('Source span does not match original fragment')
        elif r['original']!=b['raw']:
            raise ValueError('Original cells differ from source block')
        if set(column_names)!=set(r['normalized']):
            raise ValueError('Column mapping does not match normalized fields')
        if config is None:
            validate_preset(r['normalized'],p['extraction']['profile'])
        else:
            if p['extraction']['profile'] in ('cemetery_v1','parish_announcements_v1'):
                validate_preset(r['extracted'],p['extraction']['profile'])
            if set(r['field_sources'])!=set(column_names):
                raise ValueError('Field provenance does not match configured fields')
            profile = config['profile']
            if profile=='xlsx':
                from openpyxl.utils import get_column_letter
                expected_extracted = {re.sub(r'[0-9]+$','',cell):v['value'] for cell,v in r['original'].items()}
                if r['extracted'] != expected_extracted:
                    raise ValueError('Extracted spreadsheet values disagree with original cells')
            elif profile=='docx_table' and r['extracted'] != r['original']:
                raise ValueError('Extracted table values disagree with original cells')
            elif profile=='docx_paragraphs':
                match = paragraph_pattern.fullmatch(b['raw']) if paragraph_pattern else None
                expected_extracted = {'text':b['raw']}
                if match:
                    expected_extracted.update(match.groupdict())
                if r['extracted'] != expected_extracted:
                    raise ValueError('Extracted paragraph values disagree with original text')
            for field in config['fields']:
                name=field['name'];value=r['normalized'][name]
                validate_typed_value(value,field)
                if profile=='parish_announcements_v1' and field['source'].get('path')=='recurrence' and field['type']=='json' and value is not None:
                    validate_recurrence(value)
                evidence=r['field_sources'][name]
                if evidence['selector']!=field['source']:
                    raise ValueError('Field provenance selector disagrees with configuration')
                if any(evidence['location'].get(k)!=v for k,v in b['location'].items()):
                    raise ValueError('Field provenance location disagrees with source block')
                if 'path' in field['source']:
                    from .mapping import get_path
                    if evidence['value']!=get_path(r['extracted'],field['source']['path']):
                        raise ValueError('Field provenance disagrees with extracted value')
                if profile in ('xlsx','docx_table'):
                    col = bindings[name]
                    if profile=='xlsx':
                        cell = get_column_letter(col) + str(loc['row'])
                        original = r['original'][cell]
                        expected_value = original['value']
                        if field['type']=='string' and type(expected_value) is int and re.fullmatch('0+',original['number_format']):
                            expected_value = str(expected_value).zfill(len(original['number_format']))
                        expected_location = {**loc,'cell':cell}
                    else:
                        from openpyxl.utils import get_column_letter
                        expected_value = r['original'][get_column_letter(col)]
                        expected_value = None if not expected_value.strip() else expected_value
                        expected_location = {**loc,'column':col}
                    if evidence['value'] != expected_value or evidence['location'] != expected_location:
                        raise ValueError('Field provenance disagrees with mapped source cell')
                elif profile=='docx_paragraphs' and 'capture' in field['source']:
                    capture = field['source']['capture']
                    if not paragraph_pattern or capture not in paragraph_pattern.groupindex:
                        raise ValueError('Unknown paragraph capture')
                    expected_value = match[capture] if match else None
                    expected_location = dict(loc)
                    if match and match.start(capture)>=0:
                        expected_location.update(char_start=match.start(capture),char_end=match.end(capture))
                    if evidence['value'] != expected_value or evidence['location'] != expected_location:
                        raise ValueError('Field provenance disagrees with paragraph capture')
                if value is None and not field['nullable']:
                    explained=any(i['record_id']==r['id'] and i['code']=='mapping_error' and i['status']=='open'
                                  and i['severity']=='error' and i['message'].startswith(name+':') for i in p['issues'])
                    if not explained:
                        raise ValueError('Required field is missing without a blocking mapping issue')
    represented={r['block_id'] for r in p['records']}|{i['block_id'] for i in p['issues']}
    for b in p['source_blocks']:
        if b['kind'] in ('data','category','unresolved') and b['id'] not in represented:
            raise ValueError('Source content has no record or issue')
    if version=='2.1.0':
        from .decisions import fingerprint
        for entry in p['review_history']:
            if entry['revision']>p['revision'] or any(key not in issue_ids for key in entry['issue_ids']):
                raise ValueError('Invalid review history reference or revision')
        latest_decisions={}
        for entry in p['review_history']:
            if entry['action'] in ('accept_flags','resolve_flags','reopen_flags'):
                for key in entry['issue_ids']:
                    latest_decisions[key]=entry['after'].get(key)
        for finding in p['issues']:
            if finding['status']=='accepted' and latest_decisions.get(finding['id'])!='accepted':
                raise ValueError('Accepted flag is missing its review decision')
        approval=p['approval']
        if p['review_status']=='approved':
            if not approval or approval['revision']!=p['revision'] or approval['sha256']!=fingerprint(p):
                raise ValueError('Approval does not match this JSON revision')
            if not p['review_history'] or p['review_history'][-1]['action']!='approve_package':
                raise ValueError('Approval decision is missing from review history')
            entry=p['review_history'][-1]
            if any(approval[k]!=entry[k] for k in ('actor','note','at','revision')):
                raise ValueError('Approval metadata disagrees with review decision')
        elif approval is not None:
            raise ValueError('Only an approved package may carry active approval')
    if p['review_status']=='approved':
        if any(i['severity']=='error' and i['status']=='open' for i in p['issues']):
            raise ValueError('Unresolved blocking issues prevent approval')
        if any(r['review_status']=='needs_review' for r in p['records']):
            raise ValueError('All records require a review decision')
    return p
