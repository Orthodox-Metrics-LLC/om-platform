"""Cemetery profile: preserve every populated row and every original cell."""
from collections import Counter
from datetime import date, datetime
import re
from openpyxl import load_workbook
from .common import columns, issue, record, scalar

# User-approved cemetery policy: import these source conditions as written.
ACCEPTED_SOURCE_CONDITIONS = frozenset({
    "missing_name", "repeated_plot_grave", "suspicious_year", "death_before_birth"
})


def cemetery_finding(p, code, message, block_id, record_id):
    accepted = code in ACCEPTED_SOURCE_CONDITIONS
    finding = issue(p, code, message + (
        " Accepted as source data; correction is not required." if accepted else ""
    ), block_id, record_id, "warning" if accepted else "error")
    if accepted:
        finding["status"] = "resolved"
    return finding


HEADERS = ["Count", "Plot", "Grave", "First Name", "Last Name", "DOB", "DOD"]


def partial_date(value):
    if value is None:
        return None, None
    if isinstance(value, (date, datetime)):
        return {"year": value.year, "month": value.month, "day": value.day, "precision": "day"}, None
    if isinstance(value, int) and not isinstance(value, bool) and 1 <= value <= 9999:
        return {"year": value, "month": None, "day": None, "precision": "year"}, None
    m = re.fullmatch(r"\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})\s*", str(value))
    if not m:
        return None, "unresolved_date"
    month, day, year = map(int, m.groups())
    if not year or month > 12 or (day and not month):
        return None, "invalid_date"
    try:
        date(year, month or 1, day or 1)
    except ValueError:
        return None, "invalid_date"
    result = {"year": year, "month": month or None, "day": day or None,
              "precision": "day" if day else "month" if month else "year"}
    return result, "suspicious_year" if year < 1500 else None


def identifier(cell):
    value = cell.value
    if value is None:
        return None
    if isinstance(value, int) and re.fullmatch(r"0+", cell.number_format):
        return str(value).zfill(len(cell.number_format))
    return str(value).strip()


def extract_xlsx(path, p):
    workbook = load_workbook(path, data_only=False)
    cache = load_workbook(path, data_only=True)
    sheets = []
    keys = Counter()
    for sheet in workbook:
        populated = [row for row in sheet if any(c.value is not None for c in row)]
        meta = {"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                "merged_ranges": [str(x) for x in sheet.merged_cells.ranges], "blank_rows": [],
                "header_row": None}
        sheets.append(meta)
        if not populated:
            continue
        header = next((row for row in populated if [c.value for c in row[:7]] == HEADERS), None)
        if header is not None:
            meta["header_row"] = header[0].row
        for row in sheet:
            if header is not None and row[0].row <= header[0].row:
                disposition = "header" if row[0].row == header[0].row else "unresolved"
            elif not any(c.value is not None for c in row):
                meta["blank_rows"].append(row[0].row)
                continue
            else:
                disposition = "data" if header is not None else "unresolved"
            raw = {c.coordinate: {"value": scalar(c.value), "type": c.data_type,
                                  "number_format": c.number_format,
                                  "cached_value": scalar(cache[sheet.title][c.coordinate].value)} for c in row}
            b = {"id": f"sheet-{sheet.title}-row-{row[0].row}", "kind": disposition,
                 "location": {"sheet": sheet.title, "row": row[0].row}, "raw": raw}
            p["source_blocks"].append(b)
            if disposition == "header":
                continue
            if disposition == "unresolved":
                issue(p, "unsupported_sheet", "Row preserved; cemetery headers not recognized.", b['id'], severity="error")
                continue
            vals = [c.value for c in row[:7]]
            born, birth_error = partial_date(vals[5])
            died, death_error = partial_date(vals[6])
            n = {"count": scalar(vals[0]), "plot": identifier(row[1]), "grave": identifier(row[2]),
                 "first_name": str(vals[3]).strip() if vals[3] is not None else None,
                 "last_name": str(vals[4]).strip() if vals[4] is not None else None,
                 "birth": born, "death": died}
            r = record(p, b, n)
            keys[(n['plot'], n['grave'])] += 1
            for field, error in [('birth', birth_error), ('death', death_error)]:
                if error:
                    cemetery_finding(p, error, f"{field.capitalize()} value retained as written.", b['id'], r['id'])
            if not n['first_name'] or not n['last_name']:
                cemetery_finding(p, "missing_name", "Name is incomplete; row retained.", b['id'], r['id'])
            if born and died and died['year'] < born['year']:
                cemetery_finding(p, "death_before_birth", "Death year precedes birth year; values retained.", b['id'], r['id'])
            if len(row) > 7 and any(c.value is not None for c in row[7:]):
                issue(p, "extra_columns", "Additional cells preserved but not mapped.", b['id'], r['id'], "error")
            if any(c.data_type in ('f', 'e') for c in row):
                issue(p, "formula_or_cell_error", "Formula/error requires review; expression and cache retained.", b['id'], r['id'], "error")
    for r in p['records']:
        n = r['normalized']
        if keys[(n['plot'], n['grave'])] > 1:
            cemetery_finding(p, "repeated_plot_grave", "Repeated plot/grave; not deduplicated or used as a primary key.", r['block_id'], r['id'])
    columns(p, {"count": "VARCHAR(64)", "plot": "VARCHAR(128)", "grave": "VARCHAR(128)",
                "first_name": "TEXT", "last_name": "TEXT", "birth": "JSON", "death": "JSON"})
    p['extraction']['metadata'] = {"worksheets": sheets, "record_count": len(p['records']),
        "source_condition_policy": {
            "accepted_codes": sorted(ACCEPTED_SOURCE_CONDITIONS),
            "handling": "Preserve values and all rows; acknowledged notices do not require correction.",
            "decision_basis": "User explicitly accepted these four cemetery source conditions in this chat.",
            "package_approval": "Separate from accepting source conditions."
        }}
    workbook.close()
    cache.close()
