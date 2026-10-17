"""Format readers with configurable source selection and no domain field names."""
import re
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from docx import Document
from .common import record, scalar, issue
from .mapping import apply_fields


def column_bindings(config,headers):
    bindings = {}
    for f in config['fields']:
        source = f['source']
        if 'header' in source:
            matches = [i+1 for i,v in enumerate(headers) if v is not None and str(v).strip() == source['header']]
            if len(matches) != 1:
                raise ValueError(f"Header {source['header']!r} must identify exactly one column; use a column letter to disambiguate")
            col = matches[0]
        else:
            col = column_index_from_string(source['column'])
        if col > len(headers):
            raise ValueError(f"Source column for {f['name']} is outside the selected table")
        bindings[f['name']] = col
    return bindings


def extract_xlsx_generic(path,p,config):
    w = load_workbook(path,data_only=False)
    cache = load_workbook(path,data_only=True)
    try:
        sel = config['selection']
        selected = sel.get('sheet',w.sheetnames[0])
        if selected not in w.sheetnames:
            raise ValueError(f'Unknown worksheet: {selected}')
        sheet = w[selected]
        header = sel.get('header_row',1)
        start,end = sel.get('start_row',header+1),sel.get('end_row',sheet.max_row)
        if header > sheet.max_row or end > sheet.max_row:
            raise ValueError('Selected rows exceed worksheet bounds')
        if not any(c.value is not None for row in sheet for c in row):
            raise ValueError('Selected worksheet is empty')
        bindings = column_bindings(config,[c.value for c in sheet[header]])
        metadata=[]
        for s in w:
            meta={'name':s.title,'rows':s.max_row,'columns':s.max_column,
                  'selected':s.title==selected,'blank_rows':[],
                  'merged_ranges':[str(x) for x in s.merged_cells.ranges]}
            metadata.append(meta)
            for cells in s:
                number=cells[0].row
                if not any(c.value is not None for c in cells):
                    meta['blank_rows'].append(number)
                    continue
                selected_row=s.title==selected and start<=number<=end
                kind='data' if selected_row else 'excluded'
                if s.title==selected and number==header:
                    kind='header'
                raw={c.coordinate:{'value':scalar(c.value),'type':c.data_type,
                                    'number_format':c.number_format,
                                    'cached_value':scalar(cache[s.title][c.coordinate].value)} for c in cells}
                b={'id':f'sheet-{s.title}-row-{number}','kind':kind,
                   'location':{'sheet':s.title,'row':number},'raw':raw}
                p['source_blocks'].append(b)
                if selected_row:
                    r=record(p,b,{get_column_letter(c.column):scalar(c.value) for c in cells})
                    if any(c.data_type in ('f','e') for c in cells):
                        issue(p,'formula_or_cell_error','Formula/error retained; no formula evaluation.',b['id'],r['id'],'error')
                    if any(rng.min_row<=number<=rng.max_row for rng in s.merged_cells.ranges):
                        issue(p,'merged_cells','Merged data cells require review; values are not propagated.',b['id'],r['id'],'error')
        def resolver(r,f):
            col=bindings[f['name']]
            cell=get_column_letter(col)+str(r['source_location']['row'])
            original=r['original'][cell]
            value=original['value']
            error='Formula or cell error cannot be converted' if original['type'] in ('f','e') else None
            if f['type']=='string' and type(value) is int and re.fullmatch('0+',original['number_format']):
                value=str(value).zfill(len(original['number_format']))
            return value,{**r['source_location'],'cell':cell},error
        config['selection']={'sheet':selected,'header_row':header,'start_row':start,'end_row':end}
        p['extraction']['metadata']={'worksheets':metadata,'record_count':len(p['records']),
                                     'column_bindings':bindings,'selection_is_explicit':True}
        apply_fields(p,config,resolver)
    finally:
        w.close()
        cache.close()


def extract_docx_generic(path,p,config):
    d=Document(path)
    sel=config['selection']
    paragraph_mode=config['profile']=='docx_paragraphs'
    pattern=None
    if paragraph_mode and sel.get('pattern'):
        try:
            pattern=re.compile(sel['pattern'])
        except re.error as exc:
            raise ValueError(f'Invalid paragraph pattern: {exc}') from exc
        for f in config['fields']:
            if 'capture' in f['source'] and f['source']['capture'] not in pattern.groupindex:
                raise ValueError(f"Unknown named capture: {f['source']['capture']}")
    start=sel.get('start_paragraph',0)
    end=sel.get('end_paragraph',len(d.paragraphs)-1)
    if paragraph_mode and (start>=len(d.paragraphs) or end>=len(d.paragraphs)):
        raise ValueError('Selected paragraph range exceeds document bounds')
    selected_table=sel.get('table_index')
    if not paragraph_mode and selected_table>=len(d.tables):
        raise ValueError('Selected table does not exist')
    if paragraph_mode:
        config['selection'] = {'start_paragraph':start,'end_paragraph':end}
        if pattern:
            config['selection']['pattern'] = pattern.pattern
    captures={}
    for i,para in enumerate(d.paragraphs):
        text=para.text
        selected=paragraph_mode and start<=i<=end and bool(text.strip())
        b={'id':f'paragraph-{i}','kind':'data' if selected else 'excluded',
           'location':{'paragraph':i},'raw':text}
        p['source_blocks'].append(b)
        if selected:
            match=pattern.fullmatch(text) if pattern else None
            values={'text':text}
            if match:
                values.update(match.groupdict())
            r=record(p,b,values)
            captures[r['id']]=match
            if pattern and not match:
                issue(p,'unmatched_paragraph','Paragraph retained; pattern did not match the entire text.',b['id'],r['id'],'error')
    bindings={}
    for ti,table in enumerate(d.tables):
        selected=not paragraph_mode and ti==selected_table
        if selected:
            header=sel.get('header_row',1)
            first=sel.get('start_row',header+1)
            last=sel.get('end_row',len(table.rows))
            if header>len(table.rows) or last>len(table.rows):
                raise ValueError('Selected rows exceed Word table bounds')
            bindings=column_bindings(config,[c.text for c in table.rows[header-1].cells])
            config['selection']={'table_index':ti,'header_row':header,'start_row':first,'end_row':last}
        for ri,row in enumerate(table.rows,1):
            included=selected and first<=ri<=last
            values={get_column_letter(ci):c.text for ci,c in enumerate(row.cells,1)}
            b={'id':f'table-{ti}-row-{ri}','kind':'data' if included else 'excluded',
               'location':{'table':ti,'row':ri},'raw':values}
            p['source_blocks'].append(b)
            if included:
                r=record(p,b,values)
                # Merged cells are aliases of the same XML element; retain rather than duplicate silently.
                if len({id(c._tc) for c in row.cells}) < len(row.cells) or any(c._tc.xpath('./w:tcPr/w:vMerge') for c in row.cells):
                    issue(p,'merged_cells','Merged Word cells require review.',b['id'],r['id'],'error')
    def resolver(r,f):
        if paragraph_mode:
            selector=f['source']
            if 'path' in selector:
                return r['extracted']['text'],dict(r['source_location']),None
            match=captures[r['id']]
            capture=selector['capture']
            loc=dict(r['source_location'])
            if not match:
                return None,loc,'Paragraph pattern did not match'
            a,z=match.span(capture)
            if a>=0:
                loc.update(char_start=a,char_end=z)
            return match[capture],loc,None
        col=bindings[f['name']]
        value=r['extracted'].get(get_column_letter(col))
        if isinstance(value,str) and not value.strip():
            value=None
        return value,{**r['source_location'],'column':col},None
    p['extraction']['metadata']={'paragraph_count':len(d.paragraphs),'table_count':len(d.tables),
                                 'record_count':len(p['records']),'column_bindings':bindings,
                                 'selection_is_explicit':True}
    apply_fields(p,config,resolver)
