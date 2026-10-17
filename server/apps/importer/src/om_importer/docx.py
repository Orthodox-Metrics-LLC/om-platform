"""Deterministic parish prose grammar. Uncertain fragments remain review candidates."""
import calendar
from collections import Counter
import re
from docx import Document
from .common import columns, issue, record

MONTHS = {name.upper(): i for i, name in enumerate(calendar.month_name) if name}
GROUPS = {"Current Parishioners": "current_parishioners",
          "Friends of the Parish & Former Parishioners": "friends_former",
          "Fallen Asleep in the Lord": "fallen_asleep", "Archived": "archived"}
CATEGORIES = {"Namedays": "nameday", "Birthdays": "birthday",
              "Baptism Anniversaries": "baptism_anniversary",
              "Wedding Anniversaries": "wedding_anniversary",
              "Chrismation Anniversaries": "chrismation_anniversary",
              "Ordination Anniversaries": "ordination_anniversary",
              "Departed": "departed", "Slava": "slava"}
LABEL = re.compile(r'^\s*(' + '|'.join(CATEGORIES) + r')\s*[–—\-:]\s*')
ORD = r'\d{1,2}(?:st|nd|rd|th|trh)'
ANCHOR = re.compile(r'(?<![\w])(?:' + ORD + r' Sunday of Pascha|Sunday after ' + ORD
                    + r'|Sunday ' + ORD + r'\s*[-–]\s*' + ORD
                    + r'|' + ORD + r'\s*[-–]\s*' + ORD + r'|' + ORD + r')(?!\w)')
ROLES = {'baptism_anniversary': 'baptism', 'chrismation_anniversary': 'chrismation',
         'wedding_anniversary': 'wedding', 'ordination_anniversary': 'ordination', 'departed': 'death'}


def depths(text):
    depth = 0
    positions = []
    bad = False
    for c in text:
        positions.append(depth)
        if c == '(':
            depth += 1
        elif c == ')':
            if not depth:
                bad = True
            depth = max(0, depth - 1)
    return positions, bad or depth != 0


def split_spans(text, start, end, commas=True):
    ds, _ = depths(text[start:end])
    left = start
    for i in range(start, end):
        if ds[i-start] == 0 and (text[i] == ';' or (commas and text[i] == ',')):
            # Honor suffixes rather than inventing a person named Sr. or Jr.
            if text[i] == ',' and re.match(r'\s*(Sr\.?|Jr\.?)\s*(?:[,;]|$)', text[i+1:end]):
                continue
            if text[left:i].strip(' ;,.'):
                yield left, i
            left = i + 1
    if text[left:end].strip(' ;,.'):
        yield left, end


def recurrence(anchor, month):
    base = {"type": "unresolved", "text": anchor or "", "month": None, "day": None, "parameters": None}
    if not anchor or 'trh' in anchor:
        return base
    nums = [int(x) for x in re.findall(r'\d+', anchor)]
    if 'Pascha' in anchor:
        base.update(type='pascha_relative', parameters={"ordinal": nums[0], "weekday": "Sunday", "reference": "Pascha"})
    elif anchor.startswith('Sunday after'):
        base.update(type='weekday_after', parameters={"weekday": "Sunday", "after_day": nums[0], "source_month": month})
    elif anchor.startswith('Sunday'):
        base.update(type='weekday_in_range', parameters={"weekday": "Sunday", "start_day": nums[0], "end_day": nums[1], "source_month": month})
    elif len(nums) == 1 and month and 1 <= nums[0] <= calendar.monthrange(2000, month)[1]:
        base.update(type='fixed', month=month, day=nums[0])
    return base


def parse_block(p, b, event, month, group, content_start):
    text = b['raw']
    ds, bad = depths(text)
    if bad:
        issue(p, 'unbalanced_parentheses', 'Unbalanced parentheses; affected splits require review.', b['id'], severity='error')
    anchors = [m for m in ANCHOR.finditer(text, content_start) if ds[m.start()] == 0]
    segments = []
    if not anchors or text[content_start:anchors[0].start()].strip(' ;,.'):
        segments.append((content_start, anchors[0].start() if anchors else len(text), None))
    for i, a in enumerate(anchors):
        segments.append((a.end(), anchors[i+1].start() if i+1 < len(anchors) else len(text), a))
    for start, end, anchor in segments:
        rule = recurrence(anchor.group() if anchor else None, month)
        saint = None
        # Optional colon before/after saint annotation; balanced nested parentheses.
        while start < end and text[start] in ' :':
            start += 1
        if start < end and text[start] == '(':
            local_ds, _ = depths(text[start:end])
            close = next((j for j in range(1, end-start) if text[start+j] == ')' and local_ds[j] == 1), None)
            if close is not None:
                saint = text[start+1:start+close]
                start += close + 1
                while start < end and text[start] in ' :':
                    start += 1
        spans = list(split_spans(text, start, end))
        if not spans:
            spans = [(start, end)]
        ambiguous_year = (event in ('baptism_anniversary', 'chrismation_anniversary') and len(spans) > 1
                          and bool(re.search(r'\(\d{4}\)\s*[.; ]*$', text[start:end]))
                          and any(not re.search(r'\(\d{4}\)', text[a:z]) for a,z in spans))
        if ambiguous_year:
            spans = [(start, end)]
        for a, z in spans:
            subject = text[a:z].strip(' ;,.')
            year = None
            milestone = None
            subject_type = 'couple' if event == 'wedding_anniversary' else 'person'
            flags = []
            if ambiguous_year:
                subject_type = 'group'
                flags.append('ambiguous_year_scope')
            if rule['type'] == 'unresolved':
                flags.append('unresolved_date')
            if not subject:
                flags.append('missing_subject')
            if anchor and anchor.start() > content_start:
                before = text[content_start:anchor.start()].rstrip()
                if before and before[-1] not in ';,:':
                    flags.append('missing_date_delimiter')
            if bad:
                flags.append('unbalanced_parentheses')
            if not month or not group:
                flags.append('missing_context')
            if '&' in subject and event != 'wedding_anniversary':
                subject_type = 'group'
                flags.append('ambiguous_subject_split')
            if re.search(r'\)\.\s+\w|[a-z]\.\s+[A-Z]', subject) or len(subject.split()) > 6:
                flags.append('ambiguous_subject_split')
            # Preserve unseparated repeated surname patterns (e.g. Faller Carys ... Faller).
            words = re.findall(r'[A-Za-z]+', re.sub(r'\([^)]*\)', '', subject))
            if len(words) >= 4 and any(words.count(w) > 1 for w in words):
                flags.append('ambiguous_subject_split')
            milestones = [None]
            if event == 'ordination_anniversary':
                m = re.fullmatch(r'(.*?)\s*\(([^()]*)\)', subject)
                if m and all(re.fullmatch(r'\s*[^,;]+,\s*\d{4}\s*', part) for part in m[2].split(';')):
                    subject = m[1].strip()
                    milestones = [(part.rsplit(',',1)[0].strip(), int(part.rsplit(',',1)[1])) for part in m[2].split(';')]
                    if any('&' in rank or ',' in rank for rank, _ in milestones):
                        flags.append('ambiguous_milestone')
                else:
                    flags.append('ambiguous_milestone')
            elif event in ROLES and not ambiguous_year:
                m = re.fullmatch(r'(.*?)\s*\(([†+]?)(\d{4})\)', subject)
                if m:
                    subject, marker, year = m[1].strip(), m[2], int(m[3])
                    if event == 'departed' and not marker:
                        flags.append('unmarked_memorial_year')
            for milestone_item in milestones:
                if milestone_item:
                    milestone, year = milestone_item
                n = {"event_type": event, "source_group": group, "source_month": month,
                     "subject_type": subject_type, "subject_display": subject,
                     "recurrence": dict(rule), "saint_or_feast": saint,
                     "event_year": year, "year_role": ROLES.get(event) if year else None,
                     "ordination_milestone": milestone}
                r = record(p, b, n, a, z)
                for flag in sorted(set(flags)):
                    issue(p, flag, 'Review source fragment; no missing information inferred.', b['id'], r['id'], 'error')


def extract_docx(path, p):
    doc = Document(path)
    month = group = None
    category_counts, group_counts = Counter(), Counter()
    month_count = 0
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        stripped = text.strip()
        match = LABEL.match(text)
        kind = 'blank' if not stripped else 'unresolved'
        if stripped in MONTHS:
            month, group = MONTHS[stripped], None
            month_count += 1
            kind = 'month_heading'
        elif stripped in GROUPS:
            group = GROUPS[stripped]
            group_counts[group] += 1
            kind = 'group_heading'
        elif match:
            kind = 'category'
        elif i == 0:
            kind = 'title'
        elif not stripped.strip('. '):
            kind = 'punctuation'
        b = {"id": f"paragraph-{i}", "kind": kind, "location": {"paragraph": i}, "raw": text}
        p['source_blocks'].append(b)
        if match:
            event = CATEGORIES[match[1]]
            category_counts[event] += 1
            parse_block(p, b, event, month, group, match.end())
        elif kind == 'unresolved':
            issue(p, 'unclassified_text', 'Nonempty paragraph retained without a semantic mapping.', b['id'], severity='error')
    if doc.tables:
        for i, table in enumerate(doc.tables):
            b = {"id": f"table-{i}", "kind": "unresolved", "location": {"table": i},
                 "raw": [[cell.text for cell in row.cells] for row in table.rows]}
            p['source_blocks'].append(b)
            issue(p, 'unsupported_table', 'Table preserved; parish prose profile does not map tables.', b['id'], severity='error')
    columns(p, {"event_type": "VARCHAR(32)", "source_group": "VARCHAR(32)", "source_month": "TINYINT UNSIGNED",
                "subject_type": "VARCHAR(16)", "subject_display": "TEXT", "recurrence": "JSON",
                "saint_or_feast": "TEXT", "event_year": "SMALLINT UNSIGNED", "year_role": "VARCHAR(24)",
                "ordination_milestone": "TEXT"})
    p['extraction']['metadata'] = {"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables),
        "month_heading_count": month_count, "category_counts": dict(category_counts),
        "group_counts": dict(group_counts), "record_count": len(p['records']),
        "section_count": len(doc.sections), "logical_order": "OOXML body paragraphs"}
